"""Exportacao dos resultados da busca (CSV e Excel) sem carregar tudo na RAM.

- Ate EXPORTACAO_SINCRONA_MAX linhas: o arquivo e gerado em fluxo na propria
  resposta (CSV linha a linha; Excel em modo write_only num arquivo temporario).
- Acima disso: vira tarefa em segundo plano (thread do servidor), com arquivo
  em data/receita/exportacoes e link de download quando terminar.
- Colunas de contato/pessoais (telefones, e-mail, nomes de socios) so saem
  quando o usuario informa a finalidade; a exportacao fica na auditoria (so o
  numero de linhas, as colunas e a finalidade -- nunca os dados).
- Celulas que comecam com = + - @ sao neutralizadas (injecao de formula).
"""

from __future__ import annotations

import csv
import io
import logging
import tempfile
import threading
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from sqlalchemy import select

from app.config import get_settings
from app.core.formatting import csv_cell
from app.db import get_session_factory, init_database
from app.models.base_cnpj import Exportacao
from app.receita.armazenamento import exportacoes_dir
from app.services.busca_avancada import (
    COLUNAS,
    COLUNAS_PADRAO,
    Filtros,
    iterar_exportacao,
    sem_dados_de_socio,
)

log = logging.getLogger(__name__)


class ExportacaoInvalida(ValueError):
    pass


def validar_colunas(colunas: list[str] | None, finalidade: str | None) -> list[str]:
    cols = [c for c in (colunas or []) if c in COLUNAS] or list(COLUNAS_PADRAO)
    restritas = [COLUNAS[c][0] for c in cols if COLUNAS[c][2]]
    if restritas and len((finalidade or "").strip()) < 10:
        raise ExportacaoInvalida(
            "As colunas " + ", ".join(restritas) + " trazem dados de contato ou pessoais. "
            "Informe a finalidade da exportação (mínimo 10 caracteres) ou remova essas colunas."
        )
    return cols


def _valor(v: Any) -> Any:
    if isinstance(v, Decimal):
        return float(v)
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, list):
        return ",".join(str(x) for x in v)
    return v


def _celula_csv(v: Any) -> Any:
    v = _valor(v)
    if isinstance(v, float):
        return f"{v:.2f}".replace(".", ",")
    return csv_cell(v)


def gerar_csv(f: Filtros, colunas: list[str], limite: int) -> Iterator[bytes]:
    yield "﻿".encode()
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow([COLUNAS[c][0] for c in colunas])
    for i, it in enumerate(iterar_exportacao(f, colunas, limite), 1):
        w.writerow([_celula_csv(it.get(c)) for c in colunas])
        if i % 2000 == 0:
            yield buf.getvalue().encode("utf-8")
            buf.seek(0)
            buf.truncate()
    yield buf.getvalue().encode("utf-8")


def _celula_xlsx(v: Any) -> Any:
    v = _valor(v)
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def gravar_xlsx(f: Filtros, colunas: list[str], limite: int, destino: Path,
                progresso=None) -> int:
    from openpyxl import Workbook

    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Empresas")
    ws.append([COLUNAS[c][0] for c in colunas])
    n = 0
    for it in iterar_exportacao(f, colunas, limite):
        ws.append([_celula_xlsx(it.get(c)) for c in colunas])
        n += 1
        if progresso and n % 5000 == 0:
            progresso(n)
    info = wb.create_sheet("Fonte")
    info.append(["Fonte", "Receita Federal — dados abertos do CNPJ (base local)"])
    info.append(["Gerado em", datetime.now().strftime("%d/%m/%Y %H:%M")])
    info.append(["Observação", "Contatos são os declarados pela empresa à Receita e podem estar desatualizados."])
    wb.save(destino)
    return n


def xlsx_temporario(f: Filtros, colunas: list[str], limite: int) -> Path:
    tmp = tempfile.NamedTemporaryFile(prefix="busca-", suffix=".xlsx", delete=False)  # noqa: SIM115
    tmp.close()
    path = Path(tmp.name)
    gravar_xlsx(f, colunas, limite, path)
    return path


def gravar_csv(f: Filtros, colunas: list[str], limite: int, destino: Path, progresso=None) -> int:
    n = 0
    with destino.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow([COLUNAS[c][0] for c in colunas])
        for it in iterar_exportacao(f, colunas, limite):
            w.writerow([_celula_csv(it.get(c)) for c in colunas])
            n += 1
            if progresso and n % 5000 == 0:
                progresso(n)
    return n


# ---------- segundo plano ----------


def criar_exportacao(f: Filtros, colunas: list[str], formato: str, finalidade: str | None, total: int) -> dict:
    if formato not in ("csv", "xlsx"):
        raise ExportacaoInvalida("Formato inválido.")
    init_database()
    with get_session_factory()() as s:
        e = Exportacao(formato=formato, filtros=f.para_params(), colunas=colunas,
                       finalidade=(finalidade or "")[:300] or None, total=total, status="pendente")
        s.add(e)
        s.commit()
        exp_id = e.id
    threading.Thread(target=_executar, args=(exp_id,), daemon=True, name=f"exportacao-{exp_id}").start()
    return obter(exp_id)


def _executar(exp_id: int) -> None:
    factory = get_session_factory()
    with factory() as s:
        e = s.get(Exportacao, exp_id)
        if e is None:
            return
        e.status = "executando"
        s.commit()
        formato, filtros, colunas = e.formato, dict(e.filtros), list(e.colunas)
    pasta = exportacoes_dir()
    pasta.mkdir(parents=True, exist_ok=True)
    nome = f"exportacao-{exp_id}-{datetime.now().strftime('%Y%m%d%H%M%S')}.{formato}"
    destino = pasta / nome
    limite = get_settings().exportacao_max_linhas

    def progresso(n: int) -> None:
        with factory() as s2:
            x = s2.get(Exportacao, exp_id)
            if x:
                x.feitas = n
                s2.commit()

    try:
        f = Filtros.de_params(filtros)
        gravar = gravar_xlsx if formato == "xlsx" else gravar_csv
        n = gravar(f, colunas, limite, destino, progresso)
        status, erro = "concluido", None
    except Exception as ex:  # noqa: BLE001
        log.exception("exportacao %s falhou", exp_id)
        destino.unlink(missing_ok=True)
        n, status, erro = 0, "falhou", f"{type(ex).__name__}: {ex}"[:500]
    with factory() as s:
        e = s.get(Exportacao, exp_id)
        if e:
            e.status, e.erro, e.feitas = status, erro, n
            e.arquivo = nome if status == "concluido" else None
            e.finalizado_em = datetime.now()
            s.commit()


def obter(exp_id: int) -> dict | None:
    with get_session_factory()() as s:
        e = s.get(Exportacao, exp_id)
        if e is None:
            return None
        d = {c.name: getattr(e, c.name) for c in Exportacao.__table__.columns}
        d["filtros"] = sem_dados_de_socio(dict(d.get("filtros") or {}))
        return d


def listar(limite: int = 20) -> list[dict]:
    init_database()
    with get_session_factory()() as s:
        rows = s.scalars(select(Exportacao).order_by(Exportacao.id.desc()).limit(limite))
        out = []
        for e in rows:
            d = {c.name: getattr(e, c.name) for c in Exportacao.__table__.columns}
            d["filtros"] = sem_dados_de_socio(dict(d.get("filtros") or {}))  # listagem sem dado pessoal
            out.append(d)
        return out


def caminho_arquivo(exp: dict) -> Path | None:
    if not exp.get("arquivo"):
        return None
    p = exportacoes_dir() / exp["arquivo"]
    # nome vem do banco, mas confere que nao escapa da pasta
    return p if p.is_file() and p.parent.resolve() == exportacoes_dir().resolve() else None
