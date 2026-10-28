"""Paginas da base local: busca avancada, listas, socio, painel administrativo
da base e login do administrador. Funcionam sem JavaScript (formularios GET/POST
simples); o JS so acrescenta sugestoes, contadores e acompanhamento ao vivo.
"""

from __future__ import annotations

import asyncio
import io
from datetime import date
from typing import Any
from urllib.parse import urlencode

from fastapi import APIRouter, Depends, Form, Query, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response

from app.config import get_settings
from app.core.cnpj_validator import normalize_or_none
from app.receita import jobs, servico
from app.services import admin_auth
from app.services import busca_avancada as ba
from app.services import contatos as contatos_srv
from app.services import ficha as ficha_srv
from app.services import listas as listas_srv
from app.services.admin_auth import exigir_admin_pagina, ip_de
from app.services.cnae_palavras import PALAVRAS_CHAVE
from app.web.pages import UFS, _render

router = APIRouter(include_in_schema=False)


def _params(request: Request) -> dict[str, list[str]]:
    out: dict[str, list[str]] = {}
    for k, v in request.query_params.multi_items():
        if v != "":
            out.setdefault(k, []).append(v)
    return out


def _qs(params: dict[str, Any], **troca: Any) -> str:
    p = {k: v for k, v in params.items() if k not in troca}
    for k, v in troca.items():
        if v not in (None, "", []):
            p[k] = v
    return urlencode(p, doseq=True)


def _qs_sem(params: dict[str, list[str]], nome: str, valor: str | None) -> str:
    p = {k: list(v) for k, v in params.items() if k not in ("pagina",)}
    if valor is None or nome not in p:
        p.pop(nome, None)
    else:
        p[nome] = [x for x in p[nome] if x != valor]
        if not p[nome]:
            p.pop(nome)
    return urlencode(p, doseq=True)


# ---------------------------------------------------------------- busca


@router.get("/busca", response_class=HTMLResponse)
async def busca(request: Request, salva: int | None = None) -> HTMLResponse:
    params = _params(request)
    if salva and not any(k for k in params if k != "salva"):
        b = await asyncio.to_thread(listas_srv.obter_busca, salva)
        if b:
            filtros = {k: (v if isinstance(v, list) else [v]) for k, v in b["filtros"].items()}
            return RedirectResponse(f"/busca?{urlencode({**filtros, 'salva': [str(salva)]}, doseq=True)}", 303)
    params.pop("salva", None)
    ctx: dict[str, Any] = {
        "active": "busca", "params": params, "erro": None, "resultado": None, "chips": [],
        "campos": ba.CAMPOS, "grupos": ba.GRUPOS_CAMPOS, "ordenacoes": ba.ORDENACOES,
        "colunas": ba.COLUNAS, "colunas_tela": ba.COLUNAS_TELA, "colunas_padrao": ba.COLUNAS_PADRAO,
        "ufs": UFS, "atividades": PALAVRAS_CHAVE, "info_base": await asyncio.to_thread(ba.info_base),
        "salvas": await asyncio.to_thread(listas_srv.buscas_salvas),
        "recentes": [{**b, "qs": urlencode(b["filtros"], doseq=True)}
                     for b in await asyncio.to_thread(listas_srv.recentes, 8)],
        "listas": await asyncio.to_thread(listas_srv.listas),
        "municipios": [], "salva_id": salva, "qs": _qs, "qs_sem": _qs_sem,
        "combinacoes": ba.COMBINACOES_CONTATO,
        "naturezas": [], "motivos": [], "qualificacoes": [],
    }
    status = 200
    try:
        f = ba.Filtros.de_params(params)
        ctx["filtros"] = f
        nomes = await asyncio.to_thread(ba.nomes_para_chips, f)
        ctx["chips"] = ba.chips(f, nomes)
        ufs_sel = [u for u in params.get("uf", [])]
        if ufs_sel and ctx["info_base"] and ctx["info_base"]["schema_ok"]:
            ctx["municipios"] = await asyncio.to_thread(ba.lista_municipios, ufs_sel)
        if ctx["info_base"] and ctx["info_base"]["schema_ok"]:
            ctx["naturezas"] = await asyncio.to_thread(ba.lista_dominio, "naturezas")
            ctx["motivos"] = await asyncio.to_thread(ba.lista_dominio, "motivos")
            ctx["qualificacoes"] = await asyncio.to_thread(ba.lista_dominio, "qualificacoes")
        if not f.vazio() and ctx["info_base"]:
            from app.api.busca import cancelavel

            # Se o usuario cancelar (botao Parar / fechar a aba), a consulta e interrompida.
            r = await cancelavel(request, lambda ev: ba.buscar(f, cancelar=ev))
            ctx["resultado"] = r
            if salva:
                await asyncio.to_thread(listas_srv.marcar_execucao, salva, r.total)
            if f.pagina == 1:
                await asyncio.to_thread(listas_srv.registrar_recente, f.para_params(), ba.resumo_texto(f), r.total)
    except ba.FiltroInvalido as e:
        ctx["erro"], status = str(e), 400
    except ba.BuscaCancelada as e:
        ctx["erro"], status = str(e), 408
    except ba.BuscaIndisponivel as e:
        ctx["erro"], status = str(e), 503
    return _render(request, "busca.html", ctx, status)


@router.post("/busca/salvar")
async def busca_salvar(request: Request, nome: str = Form(..., max_length=120), filtros: str = Form("")) -> Response:
    from urllib.parse import parse_qs

    params = {k: v for k, v in parse_qs(filtros).items() if k not in ("pagina", "salva")}
    try:
        ba.Filtros.de_params(params)
        b = await asyncio.to_thread(listas_srv.salvar_busca, nome, params)
    except (ValueError, ba.FiltroInvalido) as e:
        return RedirectResponse(f"/busca?{filtros}&{urlencode({'erro_salvar': str(e)})}", 303)
    return RedirectResponse(f"/busca?{urlencode(params, doseq=True)}&salva={b['id']}", 303)


@router.post("/busca/salvas/{busca_id}/remover")
async def busca_remover(busca_id: int) -> Response:
    try:
        await asyncio.to_thread(listas_srv.remover_busca, busca_id)
    except LookupError:
        pass
    return RedirectResponse("/busca", 303)


# ---------------------------------------------------------------- listas


@router.get("/listas", response_class=HTMLResponse)
async def listas(request: Request) -> HTMLResponse:
    return _render(request, "listas.html", {"active": "listas", "listas": await asyncio.to_thread(listas_srv.listas)})


@router.post("/listas")
async def listas_criar(nome: str = Form(..., max_length=120), descricao: str = Form("", max_length=500)) -> Response:
    try:
        lst = await asyncio.to_thread(listas_srv.criar_lista, nome, descricao)
    except ValueError:
        return RedirectResponse("/listas", 303)
    return RedirectResponse(f"/listas/{lst['id']}", 303)


@router.get("/listas/{lista_id}", response_class=HTMLResponse)
async def lista_ver(request: Request, lista_id: int) -> HTMLResponse:
    lst = await asyncio.to_thread(listas_srv.obter_lista, lista_id)
    if lst is None:
        return _render(request, "listas.html", {"active": "listas", "listas": await asyncio.to_thread(listas_srv.listas),
                                                "erro": "Lista não encontrada."}, 404)
    atuais: dict[str, dict] = {}
    if lst["itens"] and (await asyncio.to_thread(ba.info_base) or {}).get("schema_ok"):
        f = ba.Filtros(campos={"cnpj": [i["cnpj"] for i in lst["itens"][:500]]}, por_pagina=500)
        try:
            r = await asyncio.to_thread(ba.buscar, f, contar=False)
            atuais = {i["cnpj"]: i for i in r.itens}
        except (ba.BuscaIndisponivel, ba.BuscaCancelada):
            atuais = {}
    return _render(request, "lista.html", {"active": "listas", "lista": lst, "atuais": atuais})


@router.post("/listas/adicionar")
async def listas_adicionar_form(lista_id: int = Form(...), cnpjs: list[str] = Form([]), voltar: str = Form("")) -> Response:
    try:
        await asyncio.to_thread(listas_srv.adicionar, lista_id, [{"cnpj": c} for c in cnpjs])
    except (LookupError, ValueError):
        pass
    destino = voltar if voltar.startswith("/") and not voltar.startswith("//") else f"/listas/{lista_id}"
    return RedirectResponse(destino, 303)


@router.post("/listas/{lista_id}/itens")
async def lista_adicionar(request: Request, lista_id: int, cnpjs: str = Form(""), voltar: str = Form("")) -> Response:
    itens = [{"cnpj": c} for c in cnpjs.replace("\n", ",").split(",") if c.strip()]
    try:
        await asyncio.to_thread(listas_srv.adicionar, lista_id, itens)
    except (LookupError, ValueError):
        pass
    destino = voltar if voltar.startswith("/") and not voltar.startswith("//") else f"/listas/{lista_id}"
    return RedirectResponse(destino, 303)


@router.post("/listas/{lista_id}/itens/{cnpj}/nota")
async def lista_nota(lista_id: int, cnpj: str, observacao: str = Form("", max_length=1000)) -> Response:
    try:
        await asyncio.to_thread(listas_srv.anotar, lista_id, cnpj, observacao)
    except LookupError:
        pass
    return RedirectResponse(f"/listas/{lista_id}#item-{cnpj}", 303)


@router.post("/listas/{lista_id}/itens/{cnpj}/remover")
async def lista_remover_item(lista_id: int, cnpj: str) -> Response:
    try:
        await asyncio.to_thread(listas_srv.remover_item, lista_id, cnpj)
    except LookupError:
        pass
    return RedirectResponse(f"/listas/{lista_id}", 303)


@router.post("/listas/{lista_id}/remover")
async def lista_remover(lista_id: int, confirmar: str = Form("")) -> Response:
    if confirmar == "sim":
        try:
            await asyncio.to_thread(listas_srv.remover_lista, lista_id)
        except LookupError:
            pass
        return RedirectResponse("/listas", 303)
    return RedirectResponse(f"/listas/{lista_id}", 303)


@router.get("/listas/{lista_id}/exportar.csv")
async def lista_exportar(lista_id: int) -> Response:
    import csv

    from app.core.formatting import csv_cell

    lst = await asyncio.to_thread(listas_srv.obter_lista, lista_id)
    if lst is None:
        return Response("Lista não encontrada.", status_code=404)
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["CNPJ", "Razão social (ao adicionar)", "Observação", "Adicionado em"])
    for i in lst["itens"]:
        w.writerow([i["cnpj"], csv_cell(i["razao_social"]), csv_cell(i["observacao"]),
                    i["adicionado_em"].strftime("%d/%m/%Y %H:%M") if i["adicionado_em"] else ""])
    return Response(("﻿" + buf.getvalue()).encode(), media_type="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f'attachment; filename="lista-{lista_id}.csv"'})


# ---------------------------------------------------------------- socio


@router.get("/socio/{chave}", response_class=HTMLResponse)
async def socio(request: Request, chave: str) -> HTMLResponse:
    ctx: dict[str, Any] = {"active": "socios", "socio": None, "erro": None}
    status = 200
    try:
        ctx["socio"] = await asyncio.to_thread(ficha_srv.ficha_socio, chave)
        if ctx["socio"] is None:
            ctx["erro"], status = "Sócio não encontrado na base local.", 404
    except ficha_srv.BaseIndisponivel as e:
        ctx["erro"], status = str(e), 503
    return _render(request, "socio.html", ctx, status)


# ---------------------------------------------------------------- ficha (export)


@router.get("/empresa/{cnpj}/ficha.xlsx")
async def ficha_xlsx(cnpj: str) -> Response:
    from openpyxl import Workbook

    n = normalize_or_none(cnpj)
    if n is None:
        return Response("CNPJ inválido.", status_code=400)
    try:
        f = await asyncio.to_thread(ficha_srv.ficha_local, n)
    except ficha_srv.BaseIndisponivel as e:
        return Response(str(e), status_code=503)
    if f is None:
        return Response("CNPJ não encontrado na base local.", status_code=404)
    comp = await asyncio.to_thread(contatos_srv.listar, n)

    def cel(v: Any) -> Any:
        if isinstance(v, date):
            return v.strftime("%d/%m/%Y")
        if isinstance(v, dict):
            return "Sim" if v.get("optante") else "Não"
        if isinstance(v, str) and v[:1] in "=+-@":
            return "'" + v
        return v

    wb = Workbook()
    ws = wb.active
    ws.title = "Identificação"
    ws.append(["Campo", "Valor"])
    rotulos = {
        "cnpj": "CNPJ", "matriz_filial": "Matriz/filial", "razao_social": "Razão social",
        "nome_fantasia": "Nome fantasia", "situacao": "Situação", "data_situacao": "Data da situação",
        "motivo": "Motivo", "situacao_especial": "Situação especial", "data_abertura": "Abertura",
        "natureza": "Natureza jurídica", "porte": "Porte", "capital_social": "Capital social",
        "ente_federativo": "Ente federativo responsável", "qualificacao_responsavel": "Qualificação do responsável",
        "simples": "Optante pelo Simples", "mei": "Optante pelo MEI",
    }
    for k, r in rotulos.items():
        v = f["identificacao"].get(k)
        if v not in (None, ""):
            ws.append([r, cel(v)])
    e = f["endereco"]
    for r, v in (("Logradouro", " ".join(filter(None, [e["tipo_logradouro"], e["logradouro"]]))), ("Número", e["numero"]),
                 ("Complemento", e["complemento"]), ("Bairro", e["bairro"]), ("CEP", e["cep"]),
                 ("Município", e["municipio_nome"]), ("UF", e["uf"]), ("País", e["pais_nome"])):
        if v:
            ws.append([r, cel(v)])
    ws.append(["Fonte", f"{f['base']['fonte']} — referência {f['base']['mes']}"])
    wc = wb.create_sheet("Contatos")
    wc.append(["Tipo", "Valor", "Classificação", "Origem", "Verificado em"])
    for c in f["contatos_oficiais"]:
        wc.append([c["tipo"], cel(c["formatado"]), c["subtipo"], "Oficial (Receita)", f"ref. {c['data_referencia']}"])
    for c in comp:
        wc.append([c["tipo"], cel(c["valor"]), c["subtipo"], f"Complementar: {c['fonte']}",
                   cel(c["data_verificacao"]) if c["verificado"] else "não verificado"])
    wn = wb.create_sheet("CNAEs")
    wn.append(["Tipo", "Código", "Descrição"])
    if f["cnaes"]["principal"]:
        wn.append(["Principal", f"{f['cnaes']['principal']['codigo']:07d}", f["cnaes"]["principal"]["descricao"]])
    for c in f["cnaes"]["secundarios"]:
        wn.append(["Secundária", f"{c['codigo']:07d}", c["descricao"]])
    wsoc = wb.create_sheet("Sócios")
    wsoc.append(["Nome", "Tipo", "Documento (como publicado)", "Qualificação", "Entrada", "Faixa etária",
                 "Representante legal", "Qualificação do representante"])
    for s in f["socios"]:
        rep = s["representante"] or {}
        wsoc.append([cel(s["nome"]), s["tipo"], s["documento"], s["qualificacao"], cel(s["data_entrada"]),
                     s["faixa_etaria"], cel(rep.get("nome")), rep.get("qualificacao")])
    out = io.BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="ficha-{n}.xlsx"'})


# ---------------------------------------------------------------- admin: login


def _destino_seguro(proximo: str) -> str:
    return proximo if proximo.startswith("/") and not proximo.startswith("//") and "\\" not in proximo else "/base"


@router.get("/admin/entrar", response_class=HTMLResponse)
async def admin_entrar(request: Request, proximo: str = Query("/base", max_length=300)) -> HTMLResponse:
    if admin_auth.eh_admin(request):
        return RedirectResponse(_destino_seguro(proximo), 303)
    admin_auth.token_admin()  # garante que o arquivo exista para o usuario consultar
    return _render(request, "admin_entrar.html", {
        "active": "base", "proximo": _destino_seguro(proximo), "erro": None,
        "arquivo_token": None if get_settings().admin_token else str(admin_auth._arquivo_token().resolve()),
    })


@router.post("/admin/entrar")
async def admin_entrar_post(request: Request, token: str = Form("", max_length=300),
                            proximo: str = Form("/base", max_length=300)) -> Response:
    ip = ip_de(request)
    ctx = {"active": "base", "proximo": _destino_seguro(proximo),
           "arquivo_token": None if get_settings().admin_token else str(admin_auth._arquivo_token().resolve())}
    if admin_auth.bloqueado(ip):
        jobs.auditar("admin.login", resultado="negado", detalhe="bloqueado por tentativas", ip=ip)
        return _render(request, "admin_entrar.html", {**ctx, "erro": "Muitas tentativas. Aguarde 5 minutos."}, 429)
    if not token or not admin_auth.token_confere(token):
        admin_auth.registrar_falha(ip)
        jobs.auditar("admin.login", resultado="negado", detalhe="token incorreto", ip=ip)
        return _render(request, "admin_entrar.html", {**ctx, "erro": "Token incorreto."}, 401)
    jobs.auditar("admin.login", ip=ip)
    resp = RedirectResponse(_destino_seguro(proximo), 303)
    resp.set_cookie(admin_auth.COOKIE, admin_auth.criar_sessao(), httponly=True, samesite="strict",
                    max_age=get_settings().admin_sessao_horas * 3600, secure=request.url.scheme == "https")
    return resp


@router.post("/admin/sair")
async def admin_sair(request: Request) -> Response:
    jobs.auditar("admin.logout", ip=ip_de(request))
    resp = RedirectResponse("/", 303)
    resp.delete_cookie(admin_auth.COOKIE)
    return resp


# ---------------------------------------------------------------- admin: base


@router.get("/base", response_class=HTMLResponse, dependencies=[Depends(exigir_admin_pagina)])
async def base_painel(request: Request) -> HTMLResponse:
    estado = await asyncio.to_thread(servico.estado_geral)
    job = estado["job_ativo"] or estado["ultimo_job"]
    logs = await asyncio.to_thread(jobs.logs, job["id"], 60) if job else []
    return _render(request, "base_dados.html", {
        "active": "base", "estado": estado, "job": job, "logs": logs,
        "historico": await asyncio.to_thread(jobs.historico, 20),
        "contatos": await asyncio.to_thread(contatos_srv.estatisticas),
        "settings": get_settings(),
    })


@router.post("/base/acao", dependencies=[Depends(exigir_admin_pagina)])
async def base_acao(request: Request, acao: str = Form(...), mes: str = Form(""), job_id: int = Form(0),
                    confirmar: str = Form(""), arquivo: str = Form(""), forcar: str = Form("")) -> Response:
    """Versao sem JavaScript das acoes do painel (o JS usa /api/admin/...)."""
    ip = ip_de(request)
    msg = ""
    try:
        if acao == "verificar":
            r = await asyncio.to_thread(servico.verificar_atualizacao)
            msg = r.get("erro") or f"Versão mais recente no portal: {r['mes_remoto']}"
            jobs.auditar("base.verificar", detalhe=f"mes_remoto={r.get('mes_remoto')}", ip=ip)
        elif acao == "atualizar":
            job = await asyncio.to_thread(servico.iniciar_atualizacao, mes or None, forcar=forcar == "sim")
            jobs.auditar("base.atualizar", detalhe=f"job={job['id']} mes={mes or 'mais recente'}", ip=ip)
            msg = f"Atualização iniciada (tarefa #{job['id']})."
        elif acao in ("pausar", "cancelar"):
            if acao == "cancelar" and confirmar.upper() != "CONFIRMAR":
                msg = 'Para cancelar, digite CONFIRMAR.'
            else:
                await asyncio.to_thread(jobs.solicitar, job_id, acao)
                jobs.auditar(f"base.{acao}", detalhe=f"job={job_id}", ip=ip)
                msg = "Pedido enviado; a tarefa para no próximo ponto seguro."
        elif acao == "retomar":
            await asyncio.to_thread(servico.retomar, job_id)
            jobs.auditar("base.retomar", detalhe=f"job={job_id}", ip=ip)
            msg = "Tarefa retomada de onde parou."
        elif acao == "validar":
            job = await asyncio.to_thread(servico.iniciar_validacao)
            jobs.auditar("base.validar", detalhe=f"job={job['id']}", ip=ip)
            msg = "Validação iniciada."
        elif acao == "restaurar":
            if confirmar.upper() != "CONFIRMAR":
                msg = "Para restaurar a versão anterior, digite CONFIRMAR."
            else:
                from app.receita import armazenamento as arm

                nome = await asyncio.to_thread(arm.restaurar_anterior)
                jobs.auditar("base.restaurar", detalhe=f"arquivo={nome}", ip=ip)
                msg = f"Versão ativa agora: {nome}."
        elif acao == "remover_versao":
            if confirmar.upper() != "CONFIRMAR":
                msg = "Para remover a versão, digite CONFIRMAR."
            else:
                from app.receita import armazenamento as arm

                await asyncio.to_thread(arm.remover_versao, arquivo)
                jobs.auditar("base.remover_versao", detalhe=f"arquivo={arquivo}", ip=ip)
                msg = f"Versão {arquivo} removida."
        elif acao == "limpar":
            if confirmar.upper() != "CONFIRMAR":
                msg = "Para apagar os temporários, digite CONFIRMAR."
            else:
                n = await asyncio.to_thread(servico.limpar_trabalho)
                jobs.auditar("base.limpar", detalhe=f"bytes={n}", ip=ip)
                msg = f"{n / 1024**3:.2f} GB liberados."
        elif acao == "agenda":
            form = await request.form()
            from datetime import datetime

            ativo = form.get("ativo") == "sim"
            hora = max(0, min(23, int(str(form.get("hora") or 3))))
            dia = str(form.get("dia_semana") or "")
            agenda = {"ativo": ativo, "hora": hora, "dia_semana": int(dia) if dia.isdigit() else None,
                      "ultima_execucao": datetime.now().isoformat()}
            await asyncio.to_thread(jobs.config_set, servico.CFG_AGENDA, agenda)
            jobs.auditar("base.agenda", detalhe=f"ativo={ativo} hora={hora} dia={dia or 'todos'}", ip=ip)
            msg = "Agendamento salvo."
        elif acao == "retencao":
            if confirmar.upper() != "CONFIRMAR":
                msg = "Para aplicar a retenção, digite CONFIRMAR."
            else:
                n = await asyncio.to_thread(contatos_srv.descartar_por_retencao)
                jobs.auditar("contato.retencao", detalhe=f"removidos={n}", ip=ip)
                msg = f"{n} contato(s) complementar(es) descartado(s) por retenção."
        else:
            msg = "Ação desconhecida."
    except (jobs.JobEmAndamento, ValueError, LookupError, FileNotFoundError, PermissionError, OSError) as e:
        msg = str(e) or type(e).__name__
    return RedirectResponse(f"/base?{urlencode({'msg': msg})}", 303)


@router.get("/admin/auditoria", response_class=HTMLResponse, dependencies=[Depends(exigir_admin_pagina)])
async def admin_auditoria(request: Request) -> HTMLResponse:
    return _render(request, "auditoria.html", {"active": "base", "itens": await asyncio.to_thread(jobs.auditoria, 300)})


# ---------------------------------------------------------------- admin: contatos (sem JS)


@router.post("/admin/contatos/{cnpj}", dependencies=[Depends(exigir_admin_pagina)])
async def contato_form(
    request: Request, cnpj: str, acao: str = Form("salvar"), contato_id: int = Form(0),
    tipo: str = Form(""), valor: str = Form("", max_length=320), fonte: str = Form("", max_length=200),
    url_fonte: str = Form("", max_length=500), verificado: str = Form(""), data_verificacao: str = Form(""),
    observacao: str = Form("", max_length=500),
) -> Response:
    n = normalize_or_none(cnpj)
    if n is None:
        return Response("CNPJ inválido.", status_code=400)
    ip = ip_de(request)
    erro = ""
    try:
        if acao == "remover":
            info = await asyncio.to_thread(contatos_srv.remover, n, contato_id)
            jobs.auditar("contato.remover", detalhe=f"cnpj={n} id={contato_id} tipo={info['tipo']}", ip=ip)
        else:
            oficiais = []
            try:
                f = await asyncio.to_thread(ficha_srv.ficha_local, n)
                oficiais = f["contatos_oficiais"] if f else []
            except ficha_srv.BaseIndisponivel:
                pass
            dv = date.fromisoformat(data_verificacao) if data_verificacao else None
            await asyncio.to_thread(
                contatos_srv.salvar, n, tipo=tipo, valor=valor, fonte=fonte, url_fonte=url_fonte or None,
                verificado=verificado == "sim", data_verificacao=dv, observacao=observacao,
                contato_id=contato_id or None, oficiais=oficiais,
            )
            jobs.auditar("contato.editar" if contato_id else "contato.criar",
                         detalhe=f"cnpj={n} tipo={tipo} verificado={verificado == 'sim'}", ip=ip)
    except (contatos_srv.ContatoInvalido, ValueError) as e:
        erro = str(e)
    except LookupError:
        erro = "Contato não encontrado."
    q = f"?{urlencode({'erro_contato': erro})}" if erro else ""
    return RedirectResponse(f"/empresa/{n}{q}#contatos", 303)
