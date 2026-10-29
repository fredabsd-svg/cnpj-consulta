"""Servidor MCP: ferramentas, recursos, prompts e as mesmas protecoes da interface."""

from __future__ import annotations

import json

import pytest
from mcp import Client

from app import mcp_server

CNPJ = "12345678000195"
MEI = "11222333000181"


def _dados(r):
    assert not r.is_error, r.content
    return json.loads(r.content[0].text)


@pytest.fixture
def servidor(db_initialized):
    srv = mcp_server.configurar(admin=False)
    yield srv
    mcp_server.configurar(admin=False)


async def test_lista_ferramentas_sem_admin(servidor):
    async with Client(servidor) as c:
        nomes = {t.name for t in (await c.list_tools()).tools}
    assert {"consultar_cnpj", "ficha_cnpj", "buscar_empresas", "descrever_filtros", "buscar_socio",
            "consultar_lote", "exportar_busca", "status_base_local"} <= nomes
    assert not nomes & set(mcp_server.FERRAMENTAS_ADMIN)


async def test_admin_so_com_permissao(db_initialized):
    srv = mcp_server.configurar(admin=True)
    try:
        async with Client(srv) as c:
            nomes = {t.name for t in (await c.list_tools()).tools}
        assert set(mcp_server.FERRAMENTAS_ADMIN) <= nomes
    finally:
        mcp_server.configurar(admin=False)


async def test_ferramentas_de_leitura_marcadas(servidor):
    async with Client(servidor) as c:
        tools = {t.name: t for t in (await c.list_tools()).tools}
    assert tools["buscar_empresas"].annotations.read_only_hint is True
    assert tools["consultar_cnpj"].annotations.open_world_hint is True


async def test_ficha_e_busca(base_local, servidor):
    async with Client(servidor) as c:
        f = _dados(await c.call_tool("ficha_cnpj", {"cnpj": "12.345.678/0001-95"}))
        assert f["identificacao"]["razao_social"] == "CONTABIL TESTE LTDA"
        assert {s["documento"] for s in f["socios"]} == {"***.123.456-**", "***.654.321-**"}
        b = _dados(await c.call_tool("buscar_empresas", {
            "filtros": {"uf": ["TO"], "situacao": ["2"], "atividade": ["contabilidade"],
                        "contato": "telefone_e_email"}, "por_pagina": 5}))
        assert b["total"] == 1 and b["itens"][0]["cnpj"] == CNPJ and b["mes_referencia"] == "2026-10"
        n = _dados(await c.call_tool("contar_empresas", {"filtros": {"mei": "sim"}}))
        assert n == {"total": 1}
        d = _dados(await c.call_tool("distribuicao_resultados", {"dimensoes": ["uf"]}))
        assert {x["valor"] for x in d["uf"]} == {"TO", "GO"}
        s = _dados(await c.call_tool("sugerir_empresas", {"texto": "conta"}))
        assert s["itens"][0]["cnpj"] == CNPJ
        cat = _dados(await c.call_tool("descrever_filtros", {}))
        assert any(x["nome"] == "socio_multiempresas" for x in cat["campos"])


async def test_socio(base_local, servidor):
    async with Client(servidor) as c:
        linhas = _dados(await c.call_tool("buscar_socio", {"nome": "maria silva"}))
        chave = next(x["socio_chave"] for x in linhas["itens"] if x["cnpj"] == CNPJ)
        socio = _dados(await c.call_tool("ficha_socio", {"socio_chave": chave}))
        assert socio["qtd_empresas"] == 2 and socio["identificacao_provavel"]


async def test_erros_viram_mensagem_para_o_modelo(base_local, servidor):
    async with Client(servidor) as c:
        r = await c.call_tool("ficha_cnpj", {"cnpj": "123"})
        assert r.is_error and "CNPJ inválido" in r.content[0].text
        r = await c.call_tool("buscar_empresas", {"filtros": {"uf": ["TO' OR 1=1"]}})
        assert r.is_error and "Filtro inválido" in r.content[0].text
        r = await c.call_tool("ficha_cnpj", {"cnpj": "11444777000161"})
        assert r.is_error and "não está na base local" in r.content[0].text


async def test_sem_base_local(servidor):
    async with Client(servidor) as c:
        r = await c.call_tool("buscar_empresas", {"filtros": {"uf": ["TO"]}})
    assert r.is_error and "Base de Dados CNPJ" in r.content[0].text


async def test_exportacao_exige_finalidade_e_audita(base_local, servidor):
    from openpyxl import load_workbook

    from app.receita.jobs import auditoria

    async with Client(servidor) as c:
        r = await c.call_tool("exportar_busca", {"filtros": {"uf": ["TO"]}, "colunas": ["cnpj", "email"]})
        assert r.is_error and "finalidade" in r.content[0].text
        d = _dados(await c.call_tool("exportar_busca", {
            "filtros": {"uf": ["TO"]}, "colunas": ["cnpj", "email"], "finalidade": "Prospecção do escritório"}))
    assert d["linhas"] == 3
    linhas = list(load_workbook(d["arquivo"], read_only=True)["Empresas"].iter_rows(values_only=True))
    assert linhas[0] == ("CNPJ", "E-mail") and len(linhas) == 4
    reg = next(a for a in auditoria() if a["acao"] == "busca.exportar")
    assert "origem=mcp" in reg["detalhe"] and "contato@" not in reg["detalhe"]


async def test_listas_e_buscas_salvas(base_local, servidor):
    async with Client(servidor) as c:
        lst = _dados(await c.call_tool("criar_lista", {"nome": "Prospects"}))
        assert _dados(await c.call_tool("adicionar_a_lista", {"lista_id": lst["id"], "cnpjs": [CNPJ, MEI, "x"]})) == {"adicionados": 2}
        assert len(_dados(await c.call_tool("ver_lista", {"lista_id": lst["id"]}))["itens"]) == 2
        b = _dados(await c.call_tool("salvar_busca", {"nome": "MEIs", "filtros": {"mei": "sim"}}))
        r = _dados(await c.call_tool("executar_busca_salva", {"busca_id": b["id"]}))
        assert r["total"] == 1 and r["itens"][0]["cnpj"] == MEI


async def test_consulta_local_primeiro(base_local, servidor):
    async with Client(servidor) as c:
        d = _dados(await c.call_tool("consultar_cnpj", {"cnpj": CNPJ}))
    assert d["razao_social"] == "CONTABIL TESTE LTDA"
    assert [f["fonte"] for f in d["fontes"]] == ["receita_local"]
    assert d["diligencia"]["veredicto"]["nivel"] in ("regular", "ressalvas", "critico")
    assert "campos_procedencia" not in d


async def test_recursos_e_prompts(base_local, servidor):
    async with Client(servidor) as c:
        uris = {str(r.uri) for r in (await c.list_resources()).resources}
        assert {"cnpj://busca/filtros", "cnpj://base/status"} <= uris
        rr = await c.read_resource(f"cnpj://empresa/{CNPJ}")
        assert "CONTABIL TESTE LTDA" in rr.contents[0].text
        prompts = {p.name for p in (await c.list_prompts()).prompts}
        assert {"diligencia_cadastral", "prospeccao", "mapa_de_socios"} <= prompts
        p = await c.get_prompt("prospeccao", {"perfil": "escritórios contábeis em Palmas"})
        assert "CNAE" in p.messages[0].content.text


async def test_admin_cadastra_contato_com_fonte(base_local):
    srv = mcp_server.configurar(admin=True)
    try:
        async with Client(srv) as c:
            r = await c.call_tool("cadastrar_contato_complementar",
                                  {"cnpj": CNPJ, "tipo": "email", "valor": "a@contafacil.com.br", "fonte": ""})
            assert r.is_error and "fonte" in r.content[0].text
            k = _dados(await c.call_tool("cadastrar_contato_complementar", {
                "cnpj": CNPJ, "tipo": "email", "valor": "a@contafacil.com.br", "fonte": "Site oficial"}))
            assert k["fonte"] == "Site oficial" and not k["verificado"]
            f = _dados(await c.call_tool("ficha_cnpj", {"cnpj": CNPJ}))
            assert f["contatos_complementares"][0]["valor"] == "a@contafacil.com.br"
            assert _dados(await c.call_tool("remover_contato_complementar",
                                            {"cnpj": CNPJ, "contato_id": k["id"]})) == {"removido": k["id"]}
    finally:
        mcp_server.configurar(admin=False)


def test_http_so_em_loopback():
    with pytest.raises(SystemExit, match="127.0.0.1"):
        mcp_server.main(["--http", "--host", "0.0.0.0"])


def test_admin_pela_configuracao(db_initialized, monkeypatch):
    from app.config import get_settings

    assert mcp_server.admin_permitido() is False
    monkeypatch.setenv("MCP_PERMITIR_ADMIN", "true")
    get_settings.cache_clear()
    assert mcp_server.admin_permitido() is True
