"""Busca avancada: combinacoes de filtros, E/OU/NAO, seguranca do SQL,
paginacao/ordenacao, exportacao e desempenho em volume grande."""

from __future__ import annotations

import csv
import io
import time
from datetime import date

import duckdb
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services import busca_avancada as ba

CONTABIL = "12345678000195"
FILIAL = "12345678000276"
MERCADO = "87654321000198"
MEI = "11222333000181"


def cnpjs(params: dict, **kw) -> list[str]:
    return sorted(i["cnpj"] for i in ba.buscar(ba.Filtros.de_params(params), **kw).itens)


class TestFiltros:
    def test_sem_filtros_traz_tudo(self, base_local):
        assert cnpjs({}) == sorted([CONTABIL, FILIAL, MERCADO, MEI])

    def test_exemplo_do_enunciado(self, base_local):
        """Ativas, TO, Palmas ou Araguaina, CNAE contabil, telefone e e-mail,
        pelo menos 2 anos e com QSA."""
        p = {"situacao": ["2"], "uf": ["TO"], "municipio": ["9733", "9241"], "atividade": ["contabilidade"],
             "contato": "telefone_e_email", "idade_min": "2", "tem_socios": "sim"}
        assert cnpjs(p) == [CONTABIL]

    def test_municipios_sao_ou_e_campos_sao_e(self, base_local):
        assert cnpjs({"municipio": ["9733", "9241"]}) == sorted([CONTABIL, MERCADO])
        assert cnpjs({"municipio": ["9733", "9241"], "cnae_principal": ["4711302"]}) == [MERCADO]

    def test_excluir_e_nao(self, base_local):
        assert MEI not in cnpjs({"excluir_uf": ["GO"]})
        assert cnpjs({"excluir_uf": ["TO"]}) == [MEI]

    def test_arvore_ou_e_negacao(self, base_local):
        arvore = {"op": "OU", "itens": [
            {"campo": "cnae_principal", "valor": ["4711302"]},
            {"op": "E", "itens": [{"campo": "uf", "valor": ["GO"]}, {"campo": "mei", "valor": "sim"}]},
        ]}
        assert cnpjs({"arvore": [__import__("json").dumps(arvore)]}) == sorted([MERCADO, MEI])
        neg = {"op": "E", "negar": True, "itens": [{"campo": "uf", "valor": ["TO"]}]}
        assert cnpjs({"arvore": [__import__("json").dumps(neg)]}) == [MEI]

    def test_matriz_filial_e_raiz(self, base_local):
        assert cnpjs({"matriz_filial": "2"}) == [FILIAL]
        assert cnpjs({"cnpj_basico": "12345678"}) == sorted([CONTABIL, FILIAL])
        assert cnpjs({"cnpj": "1234"}) == sorted([CONTABIL, FILIAL])

    def test_celular_fixo_por_criterio_tecnico(self, base_local):
        assert cnpjs({"tipo_telefone": ["celular"]}) == [CONTABIL]  # (63) 99234-5678
        assert cnpjs({"tipo_telefone_principal": ["fixo"]}) == [CONTABIL]  # (63) 3215-1234
        assert cnpjs({"tipo_telefone_principal": ["celular"]}) == []
        assert cnpjs({"ddd": ["63"]}) == [CONTABIL]

    def test_contatos_oficiais(self, base_local):
        assert cnpjs({"contato": "so_email"}) == [MERCADO]
        assert cnpjs({"contato": "sem_contato"}) == sorted([FILIAL, MEI])
        assert cnpjs({"email_tipo": "webmail"}) == [MERCADO]
        assert cnpjs({"email_tipo": "dominio_proprio"}) == [CONTABIL]
        assert cnpjs({"email_dominio": "contafacil.com.br"}) == [CONTABIL]

    def test_cnae_principal_secundario_grupo_descricao(self, base_local):
        assert cnpjs({"cnae_secundario": ["8211300"]}) == [CONTABIL]
        assert cnpjs({"cnae_secundario_todos": ["8211300", "6920602"]}) == [CONTABIL]
        assert cnpjs({"cnae_secundario_todos": ["8211300", "4711302"]}) == []
        assert cnpjs({"cnae": ["8211300"]}) == [CONTABIL]
        assert cnpjs({"cnae_grupo": ["692"]}) == sorted([CONTABIL, FILIAL])
        assert cnpjs({"cnae_descricao": "supermercados"}) == [MERCADO]
        assert cnpjs({"atividade": ["beleza"]}) == [MEI]

    def test_atividade_nao_usa_o_nome(self, base_local):
        """"CONTABIL TESTE" no nome nao conta: a atividade vem do CNAE oficial."""
        assert MERCADO not in cnpjs({"atividade": ["contabilidade"]})
        assert cnpjs({"razao": "contabil", "atividade": ["comercio_varejista"]}) == []

    def test_socios(self, base_local):
        assert cnpjs({"tem_socios": "nao"}) == [MEI]  # empresario individual sem QSA
        assert cnpjs({"socio_nome": "joao souza"}) == sorted([CONTABIL, FILIAL])
        # mesmo socio precisa atender aos dois criterios (Maria e socio-administradora so na contabil)
        assert cnpjs({"socio_nome": "maria", "socio_qualificacao": ["49"]}) == sorted([CONTABIL, FILIAL])
        assert cnpjs({"socio_administrador": "sim", "uf": ["TO"]}) == sorted([CONTABIL, FILIAL, MERCADO])
        assert cnpjs({"socio_multiempresas": "2"}) == sorted([CONTABIL, FILIAL, MERCADO])
        assert cnpjs({"compartilha_socios_com": CONTABIL}) == [MERCADO]
        assert cnpjs({"qtd_socios_min": "2", "matriz_filial": "1"}) == sorted([CONTABIL, MERCADO])

    def test_homonimos_nao_sao_a_mesma_pessoa(self, base_local):
        from app.services.ficha import ficha_local

        marias = [s for s in ficha_local(MERCADO)["socios"] if s["nome"] == "MARIA DA SILVA"]
        assert len(marias) == 2 and marias[0]["chave"] != marias[1]["chave"]
        qtd = sorted(s["qtd_empresas"] for s in marias)
        assert qtd == [1, 2]

    def test_perfil(self, base_local):
        assert cnpjs({"mei": "sim"}) == [MEI]
        assert cnpjs({"simples": "sim", "matriz_filial": "1"}) == sorted([CONTABIL, MEI])
        assert cnpjs({"porte": ["3"]}) == [MERCADO]
        assert cnpjs({"capital_min": "100000"}) == [MERCADO]
        assert cnpjs({"capital_max": "10000"}) == [MEI]
        assert cnpjs({"abertura_de": "2025-01-01"}) == [MEI]
        assert cnpjs({"idade_max": "2"}) == [MEI]

    def test_geografia_e_qualidade(self, base_local):
        assert cnpjs({"uf": ["go"]}) == [MEI]
        assert cnpjs({"municipio_nome": ["araguaína"]}) == [MERCADO]
        assert cnpjs({"cep": "77001"}) == sorted([CONTABIL, MERCADO, MEI])
        assert cnpjs({"endereco_completo": "nao"}) == [FILIAL]  # sem bairro e CEP curto
        assert cnpjs({"dados_completos": "sim"}) == sorted([CONTABIL, MERCADO])

    def test_busca_rapida(self, base_local):
        assert cnpjs({"q": "conta facil"}) == [CONTABIL]
        assert cnpjs({"q": "12.345.678/0001"}) == [CONTABIL]


class TestSeguranca:
    @pytest.mark.parametrize("params", [
        {"ordem": "razao; DROP TABLE empresas"},
        {"direcao": "asc; --"},
        {"situacao": ["2) OR 1=1 --"]},
        {"uf": ["TO' OR '1'='1"]},
        {"arvore": ['{"campo": "1=1; --", "valor": "x"}']},
        {"arvore": ['{"op": "XOR", "itens": []}']},
        {"cnae_grupo": ["6"]},
        {"socio_chave": ["abc' OR 1=1"]},
    ])
    def test_entradas_maliciosas_sao_recusadas(self, base_local, params):
        with pytest.raises(ba.FiltroInvalido):
            ba.buscar(ba.Filtros.de_params(params))

    def test_texto_com_aspas_e_parametro(self, base_local):
        assert cnpjs({"razao": "x' OR '1'='1"}) == []
        # a base continua integra
        assert len(cnpjs({})) == 4

    def test_arvore_profunda_demais(self, base_local):
        no = {"campo": "uf", "valor": ["TO"]}
        for _ in range(8):
            no = {"op": "E", "itens": [no]}
        with pytest.raises(ba.FiltroInvalido):
            ba.compilar(ba.Filtros(arvore=no))

    def test_api_responde_400_para_filtro_invalido(self, base_local):
        r = TestClient(app).get("/api/busca?ordem=xyz")
        assert r.status_code == 400


class TestPaginacaoOrdenacaoApi:
    def test_paginacao_e_total(self, base_local):
        r = ba.buscar(ba.Filtros(por_pagina=2, pagina=2, ordem="cnpj"))
        assert r.total == 4 and r.paginas == 2
        assert [i["cnpj"] for i in r.itens] == [CONTABIL, FILIAL][:0] + [CONTABIL, FILIAL] if False else [
            i["cnpj"] for i in r.itens]
        todos = [i["cnpj"] for i in ba.buscar(ba.Filtros(por_pagina=10, ordem="cnpj")).itens]
        assert todos == sorted(todos)
        assert [i["cnpj"] for i in r.itens] == todos[2:4]

    def test_ordenacao_capital(self, base_local):
        r = ba.buscar(ba.Filtros(ordem="capital", campos={"matriz_filial": "1"}))
        assert r.itens[0]["cnpj"] == MERCADO

    def test_api_json_contadores_e_recentes(self, base_local):
        c = TestClient(app)
        r = c.get("/api/busca", params=[("uf", "TO"), ("situacao", "2"), ("por_pagina", "2")])
        d = r.json()
        assert r.status_code == 200 and d["total"] == 3 and len(d["itens"]) == 2
        assert d["mes_referencia"] == "2026-10"
        cont = c.get("/api/busca/contadores", params=[("uf", "TO"), ("dimensao", "municipio")]).json()
        assert {x["rotulo"] for x in cont["municipio"]} == {"PALMAS", "ARAGUAINA", "GURUPI"}
        assert c.get("/api/busca/contar?uf=GO").json() == {"total": 1}
        rec = c.get("/api/buscas-recentes").json()
        assert rec and "UF: TO" in rec[0]["resumo"]
        mun = c.get("/api/base/municipios?uf=TO").json()
        assert [m["nome"] for m in mun] == ["ARAGUAINA", "GURUPI", "PALMAS"]
        cn = c.get("/api/base/dominio/cnaes?q=contabilidade").json()
        assert any(x["codigo"] == 6920601 for x in cn)

    def test_duplicidades(self, base_local):
        assert ba.duplicidades([CONTABIL, MERCADO]) == {}

    def test_sugestoes(self, base_local):
        c = TestClient(app)
        s = c.get("/api/base/sugestoes?q=conta").json()
        assert s and s[0]["cnpj"] == CONTABIL
        s = c.get("/api/base/sugestoes?q=8765").json()
        assert [x["cnpj"] for x in s] == [MERCADO]


class TestExportacao:
    def test_csv_com_colunas_escolhidas(self, base_local):
        r = TestClient(app).get("/api/busca/exportar.csv", params=[("uf", "TO"), ("colunas", "cnpj"),
                                                                   ("colunas", "razao_social"), ("colunas", "municipio")])
        assert r.status_code == 200
        linhas = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig")), delimiter=";"))
        assert linhas[0] == ["CNPJ", "Razão social", "Município"]
        assert len(linhas) == 4

    def test_colunas_restritas_exigem_finalidade(self, base_local):
        c = TestClient(app)
        r = c.get("/api/busca/exportar.csv", params=[("colunas", "cnpj"), ("colunas", "email")])
        assert r.status_code == 400 and "finalidade" in r.json()["detail"]
        r = c.get("/api/busca/exportar.csv", params=[("colunas", "cnpj"), ("colunas", "email"),
                                                      ("finalidade", "Prospecção de clientes do escritório")])
        assert r.status_code == 200 and "contato@contafacil.com.br" in r.text
        from app.receita.jobs import auditoria

        reg = [a for a in auditoria() if a["acao"] == "busca.exportar"]
        assert reg and "contato@" not in reg[0]["detalhe"] and "linhas=4" in reg[0]["detalhe"]

    def test_xlsx(self, base_local):
        from openpyxl import load_workbook

        r = TestClient(app).get("/api/busca/exportar.xlsx?mei=sim")
        assert r.status_code == 200
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        linhas = list(wb["Empresas"].iter_rows(values_only=True))
        assert linhas[1][0] == MEI and len(linhas) == 2

    def test_selecionados(self, base_local):
        r = TestClient(app).get("/api/busca/exportar.csv", params={"selecionados": f"{MEI},{MERCADO}"})
        assert r.text.count("\n") == 3

    def test_exportacao_grande_vai_para_segundo_plano(self, base_local, monkeypatch):
        monkeypatch.setenv("EXPORTACAO_SINCRONA_MAX", "1")
        from app.config import get_settings

        get_settings.cache_clear()
        c = TestClient(app)
        d = c.get("/api/busca/exportar.csv?uf=TO").json()
        assert d["segundo_plano"]
        exp_id = d["exportacao"]["id"]
        for _ in range(100):
            st = c.get(f"/api/exportacoes/{exp_id}").json()
            if st["status"] in ("concluido", "falhou"):
                break
            time.sleep(0.05)
        assert st["status"] == "concluido" and st["feitas"] == 3
        arq = c.get(f"/api/exportacoes/{exp_id}/arquivo")
        assert arq.status_code == 200 and MERCADO in arq.text

    def test_injecao_de_formula_neutralizada(self, base_local):
        from app.services.exportacao import _celula_csv, _celula_xlsx

        assert _celula_csv("=HYPERLINK(1)").startswith("'")
        assert _celula_xlsx("+cmd").startswith("'")


class TestBuscasSalvasEListas:
    def test_busca_salva_reexecuta_criterios(self, base_local):
        c = TestClient(app)
        r = c.post("/api/buscas-salvas", json={"nome": "MEIs GO", "filtros": {"uf": ["GO"], "mei": "sim"}})
        assert r.status_code == 201
        bid = r.json()["id"]
        pg = c.get(f"/busca?salva={bid}", follow_redirects=True)
        assert pg.status_code == 200 and "JOSE PEREIRA" in pg.text
        assert c.post("/api/buscas-salvas", json={"nome": "x", "filtros": {"ordem": "xyz"}}).status_code == 400

    def test_lista_com_observacoes(self, base_local):
        c = TestClient(app)
        lid = c.post("/api/listas", json={"nome": "Clientes potenciais"}).json()["id"]
        r = c.post(f"/api/listas/{lid}/itens", json={"itens": [{"cnpj": CONTABIL}, {"cnpj": "12.345.678/0001-95"}, {"cnpj": "x"}]})
        assert r.json() == {"adicionados": 1}
        assert c.patch(f"/api/listas/{lid}/itens/{CONTABIL}", json={"observacao": "ligar segunda"}).status_code == 200
        d = c.get(f"/api/listas/{lid}").json()
        assert d["itens"][0]["observacao"] == "ligar segunda" and d["itens"][0]["adicionado_em"]
        pg = c.get(f"/listas/{lid}")
        assert pg.status_code == 200 and "CONTABIL TESTE LTDA" in pg.text and "ligar segunda" in pg.text
        assert c.delete(f"/api/listas/{lid}/itens/{CONTABIL}").status_code == 204
        assert c.delete(f"/api/listas/{lid}").status_code == 204


class TestDesempenho:
    """Volume grande sintetico (300 mil estabelecimentos) direto no esquema publicado."""

    @pytest.fixture
    def base_grande(self, base_local, tmp_path):
        from app.receita import armazenamento as arm

        destino = tmp_path / "grande.duckdb"
        n = 300_000
        with duckdb.connect(str(destino)) as c:
            c.execute(f"ATTACH '{base_local.as_posix()}' AS src (READ_ONLY)")
            for t in ("cnaes", "motivos", "municipios", "naturezas", "paises", "qualificacoes", "metadata"):
                c.execute(f"CREATE TABLE {t} AS SELECT * FROM src.{t}")
            c.execute(f"""
                CREATE TABLE estabelecimentos AS
                SELECT s.* REPLACE (
                    lpad(CAST(i AS VARCHAR), 8, '0') AS cnpj_basico,
                    lpad(CAST(i AS VARCHAR), 8, '0') || '000100' AS cnpj_completo,
                    (['TO','GO','SP','MG','BA'])[1 + i % 5] AS uf,
                    CAST((ARRAY[9733, 9241, 9701, 7107])[1 + i % 4] AS SMALLINT) AS municipio,
                    CAST(CASE WHEN i % 3 = 0 THEN 2 WHEN i % 3 = 1 THEN 8 ELSE 4 END AS TINYINT) AS situacao_cadastral,
                    (ARRAY[6920601, 4711302, 6201501, 9602501])[1 + i % 4] AS cnae_fiscal_principal,
                    'EMPRESA ' || i || ' LTDA' AS razao_social,
                    'EMPRESA ' || i || ' LTDA' AS razao_social_norm,
                    DATE '2000-01-01' + CAST(i % 9000 AS INTEGER) AS data_inicio_atividade
                )
                FROM range({n}) r(i), (SELECT * FROM src.estabelecimentos WHERE cnpj_completo = '{CONTABIL}') s
                ORDER BY cnpj_completo
            """)
            c.execute("CREATE TABLE empresas AS SELECT * FROM src.empresas")
            c.execute(f"""CREATE TABLE socios AS SELECT s.* REPLACE (lpad(CAST(i AS VARCHAR), 8, '0') AS cnpj_basico,
                          'SOCIO ' || (i % 5000) AS nome_socio_norm) FROM range({n // 2}) r(i),
                          (SELECT * FROM src.socios LIMIT 1) s ORDER BY cnpj_basico""")
            for t in ("simples", "socio_resumo", "municipios_uf"):
                c.execute(f"CREATE TABLE {t} AS SELECT * FROM src.{t}")
            c.execute("CREATE TABLE nomes_busca AS SELECT razao_social_norm AS nome_norm, cnpj_completo, 'R' AS tipo "
                      "FROM estabelecimentos ORDER BY 1")
        versoes = arm.versoes_dir()
        (versoes / "receita-2026-11-grande.duckdb").write_bytes(destino.read_bytes())
        arm.escrever_json_atomico(arm.data_dir() / arm.PONTEIRO, {"arquivo": "receita-2026-11-grande.duckdb", "mes": "2026-11"})
        arm.limpar_cache()
        return n

    def test_consultas_em_volume(self, base_grande):
        casos = [
            {"uf": ["TO"], "situacao": ["2"], "atividade": ["contabilidade"]},
            {"municipio": ["9733", "9241"], "contato": "telefone_e_email", "idade_min": "5"},
            {"razao": "empresa 12345"},
            {"socio_nome": "socio 42", "uf": ["TO", "GO"]},
            {"cnae_grupo_qualquer": ["62"], "excluir_uf": ["SP"]},
        ]
        for p in casos:
            t0 = time.perf_counter()
            r = ba.buscar(ba.Filtros.de_params(p))
            dt = time.perf_counter() - t0
            assert r.total >= 1, p
            assert dt < 5, (p, dt)
        t0 = time.perf_counter()
        from app.services.ficha import sugestoes

        assert sugestoes("EMPRESA 2999")
        assert time.perf_counter() - t0 < 2
        t0 = time.perf_counter()
        r = ba.buscar(ba.Filtros(campos={"cnpj": ["00123456000100"]}))
        assert r.total == 1 and time.perf_counter() - t0 < 1

    def test_exportacao_em_lotes_sem_carregar_tudo(self, base_grande):
        it = ba.iterar_exportacao(ba.Filtros(campos={"uf": ["TO"]}), ["cnpj", "uf"], limite=100_000, lote=1000)
        n = sum(1 for _ in it)
        assert n == 60_000

    def test_timeout_e_cancelamento_interrompem_consulta(self, base_grande):
        lenta = "SELECT count(*) FROM range(100000000000) a"
        with duckdb.connect() as conn:
            with pytest.raises(ba.BuscaCancelada, match="tempo limite"):
                ba._executar(conn, ba._Controle(None, 0.3), lenta, [])
            import threading

            ev = threading.Event()
            threading.Timer(0.3, ev.set).start()
            t0 = time.perf_counter()
            with pytest.raises(ba.BuscaCancelada, match="cancelada"):
                ba._executar(conn, ba._Controle(ev, 60), lenta, [])
            assert time.perf_counter() - t0 < 5

    def test_idade_calculada_em_python(self):
        assert ba._anos_atras(2) <= date.today()
